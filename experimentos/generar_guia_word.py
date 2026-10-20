"""
Genera docs/Guia SIGNALIS.docx: la guia para el equipo, con como
abrir y usar la aplicación web, y todo el vocabulario con una
imagen de referencia por seña.

Las imagenes salen de los mismos videos con los que se entreno: una
foto por letra y por numero, y el principio y el final del movimiento
para los colores y las palabras (los dos momentos en que la mano esta
mas distinta, asi se ve por ejemplo que en APRENDER el indice toca la
frente). Volver a correrlo despues de cambiar el
vocabulario en config.py actualiza el documento.

Uso:
    python generar_guia_word.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app-python"))

import cv2
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import config
import palabras_desde_videos as de_videos
import senas_desde_videos as videos
from extraer_landmarks import ruta_crudo
from landmarks import features_de_secuencia

SALIDA = os.path.join(config.PROYECTO, "docs", "Guia SIGNALIS.docx")
BORDO = RGBColor(0x8B, 0x2E, 0x36)

# De que video sale la foto de cada categoria estatica, y que parte de
# la imagen se usa (x0, x1, y0, y1 en fracciones).
FOTOS_ESTATICAS = {
    "abecedario": ("bhzaMJw_NbE", (0.16, 0.88, 0.0, 0.92)),
    "numeros": ("gLSSENEOGD4", (0.10, 0.82, 0.0, 0.92)),
}

# Señas en las que la busqueda automatica de momentos no muestra lo que
# importa: momentos elegidos mirando el video, en fraccion del clip. En
# «aceptar» el final es el puño apoyado sobre la palma, a la altura de la
# cintura, donde la busqueda considera que la mano ya bajo.
MOMENTOS_A_MANO = {
    "050-accept": (0.33, 0.85),
}


def jpeg(imagen):
    return io.BytesIO(cv2.imencode(".jpg", imagen, [cv2.IMWRITE_JPEG_QUALITY, 85])[1].tobytes())


def recortar(frame, x0, x1, y0, y1, ancho):
    alto_f, ancho_f = frame.shape[:2]
    parte = frame[int(y0 * alto_f):int(y1 * alto_f), int(x0 * ancho_f):int(x1 * ancho_f)]
    return cv2.resize(parte, (ancho, int(ancho * parte.shape[0] / parte.shape[1])), interpolation=cv2.INTER_AREA)


def fotos_estaticas(categoria):
    """-> lista de (nombre, imagen): el frame del medio del tramo de cada seña."""
    video_id, recorte = FOTOS_ESTATICAS[categoria]
    cap = cv2.VideoCapture(os.path.join(videos.CARPETA_VIDEOS, video_id + ".mp4"))
    fotos = []
    for sena, (desde, hasta) in videos.VIDEOS[video_id]["senas"].items():
        cap.set(cv2.CAP_PROP_POS_MSEC, (desde + hasta) / 2 * 1000)
        _, frame = cap.read()
        fotos.append((sena, jpeg(recortar(frame, *recorte, 760))))
    return fotos


def momentos_clave(carpeta, video):
    """
    Los dos instantes mas distintos de la seña, en numero de frame del video.

    Un momento fijo (por ejemplo, la mitad del clip) puede caer cuando la
    mano todavia esta subiendo y no mostrar lo que distingue la seña, como
    el indice tocando la frente en «aprender». Con los puntos de las manos
    ya extraidos se buscan, entre los frames con alguna mano levantada, los
    dos que mas difieren: el principio y el final del movimiento.
    """
    d = np.load(ruta_crudo(carpeta, video))
    manos, presente = d["manos"], d["presente"]
    levantada = presente & (manos[:, :, 0, 1] < 0.75)   # muñeca por encima de la cintura
    utiles = np.flatnonzero(levantada.any(axis=1))
    if len(utiles) < 2:
        return None
    postura = (manos[utiles, :, :, :2] * levantada[utiles][:, :, None, None]).reshape(len(utiles), -1)
    distancias = np.linalg.norm(postura[:, None] - postura[None], axis=-1)
    i, j = sorted(np.unravel_index(distancias.argmax(), distancias.shape))
    paso = max(1, round(cv2.VideoCapture(os.path.join(config.CARPETA_DATASET_LSA64, carpeta, video))
                        .get(cv2.CAP_PROP_FPS) / config.FPS_OBJETIVO))
    return utiles[i] * paso, utiles[j] * paso


def momentos_en_video(video_id, desde, hasta):
    """Como momentos_clave, para una seña cortada de un video largo -> dos tiempos en segundos."""
    d = np.load(os.path.join(videos.CARPETA_CACHE, video_id + ".npz"))
    tramo = (d["tiempos"] >= desde) & (d["tiempos"] < hasta)
    feats = features_de_secuencia(d["manos"][tramo], d["presente"][tramo], d["cuerpo"][tramo],
                                  d["cuerpo_ok"][tramo], float(d["aspecto"]))
    s = feats.reshape(len(feats), 2, config.FEATURES_POR_MANO)
    # mas exigente que MANO_EN_REPOSO: en estos videos las manos descansan juntas, algo mas arriba
    levantada = (s[:, :, -1] > 0) & (s[:, :, -2] < 0.95)
    utiles = np.flatnonzero(levantada.any(axis=1))
    tiempos = d["tiempos"][tramo]
    if len(utiles) < 2:
        return desde + 0.35 * (hasta - desde), desde + 0.65 * (hasta - desde)
    postura = (d["manos"][tramo][utiles, :, :, :2] * levantada[utiles][:, :, None, None]).reshape(len(utiles), -1)
    distancias = np.linalg.norm(postura[:, None] - postura[None], axis=-1)
    i, j = sorted(np.unravel_index(distancias.argmax(), distancias.shape))
    return tiempos[utiles[i]], tiempos[utiles[j]]


def fotos_de_videos(categoria):
    """-> lista de (nombre, imagen) para los grupos que salen de un video (saludos, preguntas)."""
    fotos = []
    for carpeta in config.vocabulario(categoria):
        video_id = next(v for v, info in de_videos.VIDEOS.items() if carpeta in info["senas"])
        desde, hasta = de_videos.VIDEOS[video_id]["senas"][carpeta][0]
        cap = cv2.VideoCapture(os.path.join(videos.CARPETA_VIDEOS, video_id + ".mp4"))
        momentos = []
        for t in momentos_en_video(video_id, desde, hasta):
            cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
            _, frame = cap.read()
            momentos.append(recortar(frame, 0.18, 0.82, 0.0, 1.0, 720))
        fotos.append((config.palabra_es(carpeta), jpeg(np.hstack(momentos))))
    return fotos


def fotos_con_movimiento(categoria):
    """-> lista de (nombre, imagen): el principio y el final del movimiento, uno al lado del otro."""
    if categoria in config.DENTRO_DE:
        return fotos_de_videos(categoria)
    fotos = []
    for carpeta in config.vocabulario(categoria):
        ruta = os.path.join(config.CARPETA_DATASET_LSA64, carpeta)
        video = sorted(os.listdir(ruta))[0]
        cap = cv2.VideoCapture(os.path.join(ruta, video))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        if carpeta in MOMENTOS_A_MANO:
            cuadros = [int(total * f) for f in MOMENTOS_A_MANO[carpeta]]
        else:
            cuadros = momentos_clave(carpeta, video) or (int(total * 0.35), int(total * 0.6))
        momentos = []
        for cuadro in cuadros:
            cap.set(cv2.CAP_PROP_POS_FRAMES, int(cuadro))
            _, frame = cap.read()
            # recorte amplio: que entren las manos aunque esten lejos del cuerpo
            momentos.append(recortar(frame, 0.10, 0.84, 0.0, 1.0, 720))
        fotos.append((config.palabra_es(carpeta), jpeg(np.hstack(momentos))))
    return fotos


def titulo(doc, texto, nivel):
    parrafo = doc.add_heading(texto, level=nivel)
    for tramo in parrafo.runs:
        tramo.font.color.rgb = BORDO
        tramo.font.name = "Segoe UI"
    return parrafo


def pasos(doc, lista):
    """Lista numerada que empieza en 1 (si no, Word sigue la numeracion de la lista anterior)."""
    numeraciones = doc.part.numbering_part.numbering_definitions._numbering
    num_estilo = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstracta = numeraciones.num_having_numId(num_estilo).abstractNumId.val
    nueva = numeraciones.add_num(abstracta)
    nueva.add_lvlOverride(ilvl=0).add_startOverride(1)
    for paso in lista:
        parrafo = doc.add_paragraph(paso, style="List Number")
        numpr = parrafo._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = nueva.numId


def puntos(doc, lista):
    for punto in lista:
        doc.add_paragraph(punto, style="List Bullet")


def grilla_de_fotos(doc, fotos, columnas, ancho_cm):
    tabla = doc.add_table(rows=0, cols=columnas)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(fotos), columnas):
        celdas = tabla.add_row().cells
        for celda, (nombre, imagen) in zip(celdas, fotos[i:i + columnas]):
            parrafo = celda.paragraphs[0]
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.add_run().add_picture(imagen, width=Cm(ancho_cm))
            pie = celda.add_paragraph()
            pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tramo = pie.add_run(nombre.upper())
            tramo.bold = True
            tramo.font.size = Pt(13)
            tramo.font.color.rgb = BORDO


def escribir_guia(doc, titulo=titulo, ancho_util=17.0, compacta=False):
    """
    Escribe la guia en `doc`. La usa tambien el documento del proyecto
    (generar_documento_proyecto.py), que la lleva como anexo con su propio
    formato de titulos; ancho_util = cm entre los margenes. compacta = 3
    letras y 2 señas con movimiento por fila, para que el anexo no sea largo.
    """
    titulo(doc, "SIGNALIS", 0)
    doc.add_paragraph("Reconocimiento de Lengua de Señas Argentina. Guía del sistema y su vocabulario.")

    titulo(doc, "Qué es SIGNALIS", 1)
    doc.add_paragraph(
        "SIGNALIS es una aplicación web que mira por la cámara de la computadora y reconoce señas de la Lengua de "
        "Señas Argentina (LSA). Muestra en pantalla qué seña vio y con cuánta confianza, y tiene un modo "
        "práctica que dice si la seña que hiciste es la que querías hacer.")
    doc.add_paragraph(
        "Reconoce cuatro categorías: el abecedario, los números del 0 al 10, colores y palabras. Las "
        "palabras están en tres grupos: básicas de uso cotidiano, saludos y preguntas. Todas están en la "
        "última parte de esta guía, con una imagen de cada una.")

    titulo(doc, "Cómo abrirla", 1)
    pasos(doc, [
        "Abrí la página de SIGNALIS en el campus del instituto, con Chrome o Edge. No hay que instalar nada.",
        "Cuando el navegador pregunte, permití el uso de la cámara.",
        "La primera vez tarda un poco en cargar; después abre rápido.",
    ])
    doc.add_paragraph(
        "El reconocimiento se hace en tu propia computadora: la imagen de la cámara no se envía a ningún lado.")

    titulo(doc, "Cómo se usa", 1)
    pasos(doc, [
        "Sentate de frente a la cámara, con buena luz, de modo que se vean la cara, los hombros y las manos.",
        "Elegí una categoría en la barra de la izquierda o en «Explorar categorías». En Palabras, arriba de "
        "la lista de señas se elige el grupo: Básicas, Saludos o Preguntas.",
        "Hacé la seña. A la derecha aparece la seña detectada y la confianza.",
        "Para practicar una seña puntual, tocala en la lista que aparece debajo de las categorías: la app "
        "te dice si lo que hiciste es esa seña o no.",
    ])
    puntos(doc, [
        "Letras y números: son posturas. Mantené la mano quieta medio segundo.",
        "Colores y palabras: son movimientos. Hacé la seña completa una vez y bajá la mano.",
        "La mano que no usás puede quedar quieta a la vista, apoyada o a la altura de la cintura.",
        "Si hay más de una cámara conectada, se elige arriba a la derecha.",
    ])

    titulo(doc, "Cómo aprendió el sistema", 1)
    doc.add_paragraph(
        "El sistema aprendió las señas a partir de datos públicos: el conjunto LSA64 de la Universidad "
        "Nacional de La Plata para los colores y las palabras básicas, y videos educativos de LSA para el "
        "abecedario, los números, los saludos y las preguntas.")
    doc.add_paragraph(
        "Una parte del trabajo la hizo el equipo de estudiantes, para pulirlo: probar cada seña frente a la "
        "cámara, registrar cuáles fallaban y con eso ajustar el sistema.")

    titulo(doc, "Vocabulario", 1)
    doc.add_paragraph(
        "Las imágenes muestran a la persona tal como la ve la cámara. En los colores y las palabras, cada "
        "imagen muestra el principio y el final del movimiento, de izquierda a derecha.")

    categorias = [
        ("abecedario", "Las 27 letras. Algunas se hacen en un lugar preciso: la frente, la oreja, el mentón.", 2, 8.0),
        ("numeros", "Del 0 al 10.", 2, 8.0),
        ("colores", "Señas con movimiento.", 1, 15.5),
        ("palabras", "Palabras básicas de uso cotidiano. Señas con movimiento.", 1, 15.5),
        ("saludos", "Grupo de Palabras: saludos y fórmulas de cortesía. Señas con movimiento.", 1, 15.5),
        ("preguntas", "Grupo de Palabras: palabras para preguntar. Señas con movimiento.", 1, 15.5),
    ]
    for categoria, nota, columnas, ancho in categorias:
        if compacta:
            columnas = 3 if categoria in FOTOS_ESTATICAS else 2
            ancho = ancho_util
        ancho = min(ancho, ancho_util / columnas - 0.6)
        doc.add_page_break()
        titulo(doc, config.INFO_CATEGORIAS[categoria][0], 2)
        doc.add_paragraph(nota)
        fotos = fotos_estaticas(categoria) if categoria in FOTOS_ESTATICAS else fotos_con_movimiento(categoria)
        grilla_de_fotos(doc, fotos, columnas, ancho)

    titulo(doc, "Origen de las imágenes", 2)
    doc.add_paragraph(
        "Abecedario y números: videos «Alfabeto dactilológico LSA» y «Números en LSA» de la Dirección de "
        "Documentación. Saludos y preguntas: videos «Saludos y fórmulas de cortesía en LSA» y "
        "«Organizadores de preguntas en LSA» de la Dirección de Documentación. Colores y palabras básicas: "
        "dataset LSA64 de la Universidad Nacional de La Plata (licencia CC BY-NC-SA).")


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(11)
    for seccion in doc.sections:
        seccion.left_margin = seccion.right_margin = Cm(2)
        seccion.top_margin = seccion.bottom_margin = Cm(2)
    escribir_guia(doc)
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    doc.save(SALIDA)
    print(f"{SALIDA}: {os.path.getsize(SALIDA) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
