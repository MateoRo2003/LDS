"""
Genera docs/SIGNALIS - Documentacion del proyecto.docx: el documento del
proyecto para la Expo (descripcion, investigacion, tecnologias,
requerimientos, diseño, referencias y la guia como anexo).

Formato: A4, Times New Roman 12, interlineado doble, numeracion de
pagina desde la primera hoja de contenido. El indice es un campo de
Word: se completa al abrir el archivo en Word (o al exportarlo a PDF con
--pdf, que lo hace Word).

Uso:
    python diagramas_documento.py                 (si cambiaron los diagramas)
    python generar_documento_proyecto.py [--pdf]
"""
import argparse
import os
import subprocess
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app-python"))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config
import generar_guia_word as guia

SALIDA = os.path.join(config.PROYECTO, "docs", "SIGNALIS - Documentacion del proyecto.docx")
FIGURAS = os.path.join(config.PROYECTO, "docs", "figuras")
BORDO = RGBColor(0x8B, 0x2E, 0x36)
FUENTE = "Times New Roman"
ANCHO_UTIL = 16.0  # cm entre margenes
ROSA_GRUPO = "F2DEDF"
ROSA_FILA = "FBF3F3"


# ---------------------------------------------------------------------
# Formato
# ---------------------------------------------------------------------

def fuente(estilo_o_run, tam=None, negrita=None, color=None):
    f = estilo_o_run.font
    f.name = FUENTE
    rpr = estilo_o_run.element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.insert(0, fuentes)
    for clave in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(clave), FUENTE)
    for tema in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fuentes.attrib.pop(qn(tema), None)
    if tam:
        f.size = Pt(tam)
    if negrita is not None:
        f.bold = negrita
    if color is not None:
        f.color.rgb = color


def preparar_estilos(doc):
    normal = doc.styles["Normal"]
    fuente(normal, 12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    for nombre, centrado in (("Heading 1", True), ("Heading 2", False)):
        estilo = doc.styles[nombre]
        fuente(estilo, 12, negrita=True, color=RGBColor(0, 0, 0))
        estilo.font.italic = False
        estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
        estilo.paragraph_format.line_spacing = 2.0
        estilo.paragraph_format.space_before = Pt(0)
        estilo.paragraph_format.space_after = Pt(0)
        estilo.paragraph_format.keep_with_next = True
    for nombre in ("List Bullet", "List Number", "Title"):
        fuente(doc.styles[nombre])


def cuerpo(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.27)
    agregar_texto(p, texto)
    return p


def agregar_texto(p, texto):
    """Texto con *cursiva* marcada entre asteriscos."""
    for i, tramo in enumerate(texto.split("*")):
        if tramo:
            r = p.add_run(tramo)
            r.italic = i % 2 == 1


def vinetas(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        agregar_texto(p, item)


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    tcpr.append(sombra)


def bordes_tabla(tabla, color="E6CFD1"):
    tblpr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single" if lado in ("top", "bottom", "insideH") else "nil")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), color)
        bordes.append(b)
    tblpr.append(bordes)


def texto_celda(celda, texto, tam=10, negrita=False, color=None):
    p = celda.paragraphs[0]
    r = p.add_run(texto)
    fuente(r, tam, negrita=negrita, color=color)


def anchos(tabla, cms):
    tabla.autofit = False
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, cms):
            celda.width = Cm(ancho)


def campo(parrafo, instruccion):
    """Un campo de Word (numero de pagina, indice) que Word calcula al abrir."""
    r = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, ""), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
            if tipo == "begin":
                el.set(qn("w:dirty"), "true")
        elif texto:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        else:
            el = OxmlElement("w:t")
            el.text = " "
        r._r.append(el)


def figura(doc, nombre, ancho_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(os.path.join(FIGURAS, nombre), width=Cm(ancho_cm))


def salto_de_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def compactar_tablas(doc):
    """En las tablas: interlineado simple y sin sangria."""
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------------
# Partes del documento
# ---------------------------------------------------------------------

def portada(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(os.path.join(FIGURAS, "logo_itec.jpg"), width=Cm(5.2))
    for texto, tam, color in (("EXPO INFORMÁTICA", 12, None), ("SIGNALIS", 18, BORDO)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if color is None:
            p.paragraph_format.space_before = Pt(18)
        r = p.add_run(texto)
        fuente(r, tam, negrita=True, color=color)
        r.italic = True
    p = doc.add_paragraph("Reconocimiento de Lengua de Señas Argentina")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("Curso y Orientación: 5to “A” Informática.")
    doc.add_paragraph("Profesor: Rojas Mateo.")
    alumnas = ["Alumnas: Benitez Mia, Guadalupe", "Benitez Paula, Victoria", "Nestto Nataly, Nicol",
               "Villalba Sonia, Aylen", "Villalva Florencia, Yasmin"]
    for i, nombre in enumerate(alumnas):
        p = doc.add_paragraph(nombre)
        p.paragraph_format.line_spacing = 1.15
        if i == 0:
            p.paragraph_format.space_before = Pt(8)


def indice(doc):
    salto_de_pagina(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(p.add_run("Índice"), 12, negrita=True)
    campo(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')


def seccion_con_numeros(doc):
    """Desde aca las paginas van numeradas, empezando en 1, arriba a la derecha."""
    seccion = doc.add_section(WD_SECTION.NEW_PAGE)
    seccion.header.is_linked_to_previous = False
    p = seccion.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo(p, "PAGE")
    numeracion = OxmlElement("w:pgNumType")
    numeracion.set(qn("w:start"), "1")
    seccion._sectPr.append(numeracion)


def descripcion(doc):
    doc.add_heading("Descripción del proyecto", 1)
    doc.add_heading("Problema o necesidad detectada", 2)
    cuerpo(doc, "Las lenguas de señas son lenguas naturales creadas por las comunidades sordas y transmitidas de "
           "generación en generación; no son códigos artificiales ni versiones manuales de las lenguas orales "
           "(Universidad de Buenos Aires [UBA], s.f.). Poseen una gramática propia, completa y compleja, y se "
           "caracterizan por ser visuales, espaciales, manuales y gestuales (Agencia Nacional de Discapacidad "
           "[ANDIS], s.f.). En la Argentina, esa lengua es la Lengua de Señas Argentina (LSA), que desde 2023 "
           "cuenta con reconocimiento legal como lengua natural y originaria de las personas sordas (Ley 27.710, "
           "2023).")
    cuerpo(doc, "Aprender LSA requiere practicar con el cuerpo y con la vista. Cada seña se describe a partir de la "
           "configuración de la mano, su posición y su movimiento (Ronchetti et al., 2016), y quien aprende "
           "necesita comprobar si lo que hace coincide con la seña que quiere producir. Los recursos de consulta "
           "más comunes, como los diccionarios, las láminas o los videos, muestran cómo se hace una seña, pero no "
           "devuelven ninguna información sobre lo que hace la persona que practica. Por eso, fuera de un curso "
           "con docente, resulta difícil saber si una seña está bien hecha.")
    cuerpo(doc, "A esto se suma que no existe una lengua de señas universal: la Federación Mundial de Sordos estima "
           "que se utilizan más de 300 lenguas de señas en el mundo (Servicio de Información sobre Discapacidad "
           "[SID-INICO], s.f.). Por lo tanto, las herramientas desarrolladas para otras lenguas no resuelven la "
           "necesidad de quien quiere aprender LSA, y entrenar un sistema de reconocimiento para una lengua nueva "
           "requiere contar con un conjunto de datos propio de esa lengua (Ronchetti et al., 2016).")

    doc.add_heading("Propuesta de solución", 2)
    cuerpo(doc, "SIGNALIS es una aplicación web que, a través de la cámara de la computadora, reconoce señas de la "
           "LSA y muestra en pantalla qué seña detectó y con cuánta confianza. Además, ofrece un modo práctica: la "
           "persona elige una seña de la lista, la realiza frente a la cámara y la aplicación le indica si lo que "
           "hizo corresponde a esa seña o no.")
    cuerpo(doc, "El sistema reconoce 88 señas organizadas en cuatro categorías: el abecedario (27 letras), los "
           "números del 0 al 10, ocho señas de colores (opaco, rojo, verde, amarillo, claro, celeste, colores y "
           "rosa) y palabras. Las palabras se ordenan en tres grupos: 22 palabras básicas de uso cotidiano, como "
           "gracias, ayuda, agua o comida; 13 saludos y fórmulas de cortesía, como hola, chau, buen día o por "
           "favor; y 9 palabras para hacer preguntas, como cuándo, cómo, quién o por qué. Gracias y dónde "
           "aparecen en dos grupos. Las letras y los números son posturas, que se reconocen cuando la mano "
           "permanece quieta alrededor de medio segundo; los colores y las palabras son señas con movimiento, que "
           "se reconocen a partir de la seña completa.")
    cuerpo(doc, "La aplicación funciona desde el navegador (Chrome o Edge), sin instalar nada, y el reconocimiento "
           "se realiza en la propia computadora de la persona: la imagen de la cámara no se envía a ningún "
           "servidor. El proyecto incluye además una guía de uso con el vocabulario completo y una imagen de cada "
           "seña, que se adjunta como anexo de este documento.")

    doc.add_heading("Usuarios y alcance", 2)
    cuerpo(doc, "El sistema contempla un único tipo de usuario: la persona que usa la aplicación para conocer, "
           "practicar o comprobar señas de la LSA. Puede tratarse de estudiantes que cursan LSA, docentes que la "
           "utilizan como apoyo en clase, familiares o personas interesadas en aprender. No se requiere registro "
           "ni inicio de sesión, y no existe un rol de administración dentro de la aplicación: el vocabulario se "
           "actualiza modificando los modelos y los archivos del proyecto. Fuera de la aplicación, el equipo de "
           "estudiantes cumple el rol de quien prepara los datos, prueba cada seña frente a la cámara y ajusta el "
           "sistema, tal como se describe en la sección Tecnologías y lenguajes utilizados.")
    cuerpo(doc, "Quedan fuera del alcance de esta versión la traducción de frases, oraciones o conversaciones "
           "completas, ya que el sistema reconoce señas de a una; el reconocimiento de señas que no pertenezcan "
           "a las cuatro categorías mencionadas o de otras lenguas de señas distintas de la LSA; y las cuentas "
           "personales y el seguimiento del progreso entre sesiones, porque el historial de señas reconocidas "
           "solo se conserva mientras la página permanece abierta. Tampoco se contempla una aplicación móvil "
           "nativa o instalable, dado que la propuesta se limita a una aplicación web pensada para computadora, "
           "ni el envío o el almacenamiento de las imágenes de la cámara en servidores.")


def objetivos(doc):
    doc.add_heading("Objetivos", 1)
    doc.add_heading("Objetivo general", 2)
    cuerpo(doc, "Desarrollar una aplicación web que reconozca señas de la Lengua de Señas Argentina mediante la "
           "cámara de la computadora y brinde una devolución inmediata a quien practica, sin requerir instalación "
           "y sin enviar las imágenes a servidores.")
    doc.add_heading("Objetivos específicos", 2)
    vinetas(doc, [
        "Investigar la problemática del aprendizaje de la LSA y los antecedentes de herramientas similares.",
        "Definir un vocabulario inicial organizado en cuatro categorías (abecedario, números, colores y "
        "palabras) y ordenar las palabras en tres grupos: básicas, saludos y preguntas.",
        "Reunir y organizar datos públicos de LSA y entrenar con ellos los modelos que reconocen las señas.",
        "Implementar la detección de las manos en el navegador y un motor de reconocimiento, basado en redes "
        "neuronales, que distinga posturas y movimientos.",
        "Mostrar la seña detectada junto con el nivel de confianza e incorporar un modo práctica que indique si "
        "la seña realizada es la esperada.",
        "Probar cada seña frente a la cámara, registrar cuáles fallan y ajustar el sistema.",
        "Publicar la aplicación como sitio web accesible desde el navegador, en el campus del instituto, y "
        "elaborar una guía de uso con el vocabulario completo.",
    ])


def investigacion(doc):
    doc.add_heading("Investigación de la problemática", 1)
    cuerpo(doc, "Este proyecto no incluyó entrevistas ni encuestas a usuarios. La problemática se abordó mediante "
           "una investigación documental y exploratoria, es decir, a partir de la lectura y el análisis de "
           "fuentes ya publicadas. La búsqueda se organizó en cuatro líneas: la normativa y los documentos "
           "institucionales sobre la LSA; los artículos académicos y los conjuntos de datos de señas argentinas; "
           "la documentación técnica de las herramientas de visión por computadora disponibles para navegadores; "
           "y los proyectos similares publicados en repositorios abiertos.")
    cuerpo(doc, "Según la Federación Mundial de Sordos, hay en el mundo aproximadamente 70 millones de personas "
           "sordas, que utilizan más de 300 lenguas de señas diferentes (SID-INICO, s.f.). Las lenguas de señas "
           "difieren entre países e incluso entre regiones (Ronchetti et al., 2016), por lo que la LSA es la "
           "lengua de la comunidad sorda argentina y no una traducción de ninguna otra. En el ámbito educativo, "
           "la UBA dicta cursos de LSA a cargo íntegramente de docentes sordos, organizados en siete niveles "
           "cuatrimestrales (UBA, s.f.). Esto muestra que el aprendizaje de la LSA se apoya en la interacción "
           "directa con docentes sordos, algo que una herramienta digital puede complementar pero no reemplazar.")
    cuerpo(doc, "En cuanto al marco normativo, la Ley 27.710, sancionada por el Congreso de la Nación en abril de "
           "2023, tiene por objeto reconocer a la LSA como lengua natural y originaria de las personas sordas en "
           "todo el territorio nacional (Ley 27.710, 2023). Su artículo 4, dedicado a la promoción de la LSA, "
           "dispone fomentar e impulsar el acceso y el uso de esta lengua por parte de todas las personas que, por "
           "cualquier motivo, elijan comunicarse en ella. Este marco legal respalda el desarrollo de recursos que "
           "acerquen la LSA a más personas, como el que propone este proyecto.")
    cuerpo(doc, "A partir de la información relevada se identificaron cuatro problemáticas, que dieron origen a los "
           "requerimientos del sistema. La primera es la falta de retroalimentación al practicar: los recursos de "
           "consulta muestran cómo hacer una seña, pero no indican si quien practica la realiza correctamente. La "
           "segunda es la especificidad de la lengua: las herramientas y los modelos creados para otras lenguas "
           "de señas no sirven para la LSA, y entrenar un reconocedor exige datos propios de esa lengua "
           "(Ronchetti et al., 2016). La tercera es la privacidad: un sistema que analiza la imagen de una persona "
           "debería poder funcionar sin enviar esa imagen a servidores, y existen librerías que procesan las "
           "imágenes en el propio dispositivo (Google, s.f.). La cuarta son las barreras de acceso: para que una "
           "herramienta de práctica llegue a más personas, no debería exigir instalar programas ni crear cuentas.")

    doc.add_heading("Antecedentes y soluciones existentes", 2)
    cuerpo(doc, "Se revisaron proyectos públicos relacionados con la LSA. Impulsar (LSAorg, s.f.) es una "
           "plataforma de software libre que recibe audio y genera su traducción a LSA mediante animaciones "
           "tridimensionales; su demostración reconoce tres instrucciones: «Hola», «Sí» y «No». Trabaja, por lo "
           "tanto, en sentido inverso al de SIGNALIS, ya que va de la voz a las señas. Entreclases (Lucho-code, "
           "s.f.) es una aplicación web para practicar LSA entre una clase y otra; su primera versión está pensada "
           "para un piloto de dos semanas con un grupo de un curso de LSA, por lo que se orienta a un grupo en "
           "particular, mientras que SIGNALIS es de uso abierto, sin registro, y se centra en el reconocimiento "
           "con cámara.")
    cuerpo(doc, "Por último, lsa-sign-language-to-speech (Mohammedalmohawis, s.f.) es un sistema que convierte "
           "señas de LSA en voz: usa MediaPipe, clasificadores temporales (GRU y Transformer), un servidor "
           "FastAPI, una interfaz en React y síntesis de voz sin conexión, y se entrena con LSA64. Su enfoque es "
           "el más cercano al de SIGNALIS, pero requiere instalar y ejecutar un servidor en Python, mientras que "
           "SIGNALIS se ejecuta completo en el navegador, sin servidor de reconocimiento.")

    doc.add_heading("Conclusiones de la investigación", 2)
    cuerpo(doc, "La investigación permitió concluir que existe una necesidad concreta de herramientas de práctica "
           "específicas para la LSA, que den una devolución inmediata a quien aprende. También mostró que hay "
           "datos públicos de LSA y librerías abiertas que permiten construir una solución de bajo costo que "
           "funcione en el navegador. A partir de estas conclusiones se definieron el alcance del proyecto, el "
           "vocabulario inicial de cuatro categorías y las decisiones tecnológicas que se justifican en la "
           "sección siguiente.")


def tecnologias(doc):
    doc.add_heading("Tecnologías y lenguajes utilizados", 1)
    cuerpo(doc, "Esta sección describe con qué lenguajes y herramientas se desarrolló SIGNALIS y por qué se "
           "eligieron. Cada decisión responde a los problemas identificados en la investigación y a cinco "
           "criterios: que la aplicación se abra desde el navegador sin instalar nada; que la imagen de la cámara "
           "se procese en la computadora de la persona; que sea de bajo costo, para lo cual se priorizaron "
           "herramientas libres o gratuitas; que sea simple de mantener, con pocos archivos y sin pasos de "
           "compilación; y que existan datos públicos de LSA con los que el sistema pueda aprender. El proyecto "
           "completo consta de 27 archivos y ocupa alrededor de 3,6 MB.")

    doc.add_heading("Lenguajes de la aplicación: HTML5, CSS3 y JavaScript", 2)
    cuerpo(doc, "La pantalla de SIGNALIS está escrita en HTML5 (index.html) y se presenta con CSS3 (style.css). Se "
           "eligieron porque son los estándares de la web: cualquier navegador los interpreta sin instalar nada "
           "y, al ser lenguajes de marcado y de estilos, no requieren herramientas de compilación. Con CSS se "
           "resolvieron además dos aspectos de usabilidad: un modo claro y un modo oscuro, definidos mediante "
           "variables de color que cambian con un solo ajuste, y una disposición que se adapta a pantallas más "
           "angostas.")
    cuerpo(doc, "El comportamiento está programado en JavaScript, que es el lenguaje que los navegadores ejecutan "
           "de forma nativa. Esto permite analizar la imagen de la cámara en la propia computadora y en tiempo "
           "real, sin necesidad de un servidor. El código se organizó en dos módulos: app.js, que maneja la "
           "cámara, los botones, el historial, los ajustes y el modo práctica, y motor.js, que realiza el "
           "reconocimiento. El motor no sabe nada de cámaras ni de pantallas: recibe los puntos de las manos y "
           "del cuerpo de cada instante y devuelve qué seña se hizo. Esta separación cumple el requerimiento de "
           "modularidad y permite probar el motor de forma independiente de la interfaz; de hecho, app.js incluye "
           "un modo de prueba que reemplaza la cámara por un video grabado, lo que facilita repetir las pruebas "
           "con los mismos movimientos.")
    cuerpo(doc, "La aplicación aprovecha además varias funciones que ya trae el navegador: el acceso a la cámara, "
           "el dibujo de los puntos de la mano sobre la imagen, la descarga de los modelos y el almacenamiento "
           "local, donde se recuerdan la cámara elegida y el tema claro u oscuro, sin guardar ningún dato "
           "personal. Se trabajó sin frameworks, en lo que el entorno del campus clasifica como «HTML5 · Vanilla "
           "Web». Como se trata de una sola pantalla, el código propio ocupa cuatro archivos (unas 2.000 líneas) "
           "y no hay pasos de compilación que puedan fallar. La contrapartida es que ciertas tareas, como "
           "actualizar la pantalla, se programan a mano, pero para una aplicación de estas dimensiones esa "
           "complejidad es manejable.")

    doc.add_heading("Detección de manos y cuerpo: MediaPipe Tasks Vision", 2)
    cuerpo(doc, "Para ubicar a la persona en la imagen se utilizó MediaPipe Tasks Vision, una librería abierta de "
           "Google (Google, s.f.), de la que el proyecto incluye el archivo vision_bundle.js en la carpeta "
           "vendor/tasks-vision. Se usan dos de sus modelos. El detector de manos ubica hasta dos manos y entrega "
           "21 puntos por cada una, que corresponden a las articulaciones y a las puntas de los dedos. El "
           "detector de pose, en su versión liviana, ubica la nariz y los dos hombros. Este segundo modelo es "
           "necesario porque muchas señas dependen del lugar donde se hacen, como la frente, la oreja o el "
           "mentón: para saberlo, el sistema mide la posición de la mano respecto de la cara y en unidades de "
           "ancho de hombros, de modo que el resultado no dependa de qué tan lejos de la cámara se siente la "
           "persona.")
    cuerpo(doc, "MediaPipe se ejecuta en el dispositivo, por lo que los datos de la imagen no se envían a "
           "servidores (Google, s.f.), lo que permite cumplir el requisito de privacidad. La aplicación intenta "
           "usar primero la placa de video, que es más rápida, y si no puede recurre al procesador. Además, la "
           "librería resuelve por sí sola la detección de las manos, de modo que el trabajo del equipo pudo "
           "concentrarse en reconocer cada seña. Otros proyectos de reconocimiento de LSA también combinan "
           "MediaPipe con el conjunto de datos LSA64 (Mohammedalmohawis, s.f.), lo que respalda la elección.")
    cuerpo(doc, "Además de la librería, MediaPipe necesita un motor de cálculo (WebAssembly) y los dos modelos "
           "mencionados. La aplicación los toma de la propia carpeta del sitio si están disponibles y, si no, los "
           "descarga de direcciones públicas (jsDelivr y el almacenamiento de modelos de Google). Por ese motivo "
           "la primera carga requiere conexión a internet y puede demorar. Lo que viaja por la red son esos "
           "archivos hacia la computadora de la persona: la imagen de la cámara nunca sale de ella.")

    doc.add_heading("Reconocimiento de señas: redes neuronales en JavaScript", 2)
    cuerpo(doc, "El reconocimiento propiamente dicho lo realizan redes neuronales, una por categoría o grupo, que "
           "reciben los puntos de las manos y devuelven la probabilidad de cada seña. Antes de entrar a la red, "
           "los puntos se transforman: cada punto de la mano se mide desde la muñeca y se divide por el tamaño de "
           "la mano, y su ubicación se expresa respecto de la cara. Así, la forma de la mano no depende de la "
           "distancia a la cámara. Si la mano principal es la izquierda, la imagen se espeja, de modo que una "
           "misma seña valga con cualquiera de las dos manos.")
    cuerpo(doc, "Para las letras y los números, que son posturas, alcanza con una sola imagen. La red recibe 135 "
           "números por imagen y es una red densa de tres capas, con 128 y 64 neuronas y una salida que da una "
           "probabilidad por seña (27 clases en el abecedario y 11 en los números), con unos 27.000 parámetros "
           "cada una. Para evitar falsos reconocimientos, el sistema promedia las últimas ocho predicciones y "
           "acepta una seña recién cuando se sostiene alrededor de medio segundo con una confianza superior al "
           "mínimo elegido.")
    cuerpo(doc, "Los colores y las palabras, en cambio, son movimientos y no pueden entenderse con una sola imagen. "
           "Para ellos el sistema arma una secuencia de 30 imágenes, con 132 números cada una, que se procesa con "
           "redes recurrentes de tipo LSTM (Hochreiter y Schmidhuber, 1997), capaces de recordar lo ocurrido en "
           "los instantes anteriores. Se usaron dos capas bidireccionales de 64 y 32 unidades (Schuster y "
           "Paliwal, 1997), que leen la secuencia hacia adelante y hacia atrás, seguidas de una capa densa de 64 "
           "neuronas y de la salida de clasificación, con unos 147.000 parámetros por modelo. Las palabras usan "
           "tres modelos, uno por grupo (básicas, saludos y preguntas), de modo que cada uno distingue entre "
           "pocas señas. Estos modelos incluyen una clase «nada» para los momentos en que la persona no hace "
           "ninguna seña. Como nadie hace una seña siempre a la misma velocidad, el sistema prueba con tres "
           "ventanas de tiempo (1; 1,5 y 2,2 segundos), se queda con la que el modelo ve más clara, promedia tres "
           "predicciones y exige que una seña nueva se sostenga durante dos predicciones seguidas.")
    cuerpo(doc, "Los modelos se entrenan en Python y sus pesos se exportan a archivos JSON, con el orden de capas y "
           "compuertas de Keras (Chollet, 2015). En el navegador, motor.js repite las cuentas de cada capa (capas "
           "densas, LSTM y salida de probabilidades) sin utilizar ninguna librería de aprendizaje automático. Se "
           "eligió esta alternativa porque las redes son pequeñas (los seis modelos juntos pesan alrededor de "
           "3,4 MB), porque evita sumar una dependencia pesada y porque permite comprobar que el motor da los "
           "mismos resultados que la versión de escritorio en Python con la que se entrenó el sistema. Se eligió "
           "JSON porque es un formato de texto que JavaScript lee de forma nativa y que se entrega como cualquier "
           "otro archivo de un sitio estático; su costo es que ocupa más espacio que un formato binario, algo "
           "aceptable para modelos tan pequeños.")
    cuerpo(doc, "Los modelos de colores, palabras, saludos y preguntas, por su tamaño, se dividen en cuatro partes "
           "de aproximadamente 250 KB cada una, porque hay sitios que no aceptan archivos grandes. El archivo "
           "principal de cada modelo indica en cuántas partes está dividido, y la aplicación las descarga y las "
           "une. Si falta algún modelo, la categoría o el grupo correspondiente queda fuera de servicio y el "
           "resto de la aplicación sigue funcionando.")

    doc.add_heading("Datos y entrenamiento", 2)
    cuerpo(doc, "Las señas de colores y las palabras básicas se aprendieron del conjunto de datos LSA64, de la "
           "Universidad Nacional de La Plata, que contiene 3.200 videos de 64 señas de la LSA realizadas por 10 "
           "personas, con cinco repeticiones cada una (Ronchetti et al., 2016). El abecedario se aprendió de tres "
           "videos educativos y los números de dos, entre ellos «Alfabeto dactilológico LSA» y «Números en LSA», "
           "de la Dirección de Documentación (Dirección de Documentación, s.f.-a, s.f.-b; Universidad Popular, "
           "s.f.; chany3h, s.f.; Sarria, s.f.). Los saludos y las preguntas se aprendieron de otros dos videos de "
           "la misma serie: «Saludos y fórmulas de cortesía en LSA» y «Organizadores de preguntas en LSA» "
           "(Dirección de Documentación, s.f.-c, s.f.-d). Se eligieron datos públicos y específicos de la LSA "
           "porque evitan tener que grabar un corpus propio completo y porque LSA64 fue creado para el "
           "reconocimiento automático de señas.")
    cuerpo(doc, "Estos datos tienen condiciones que el proyecto debió tener en cuenta. LSA64 se distribuye con "
           "licencia CC BY-NC-SA, por lo que el proyecto se plantea con fines educativos y no comerciales, e "
           "indica el origen de los datos. Además, sus videos fueron grabados con personas que usaban guantes de "
           "colores para facilitar el seguimiento de las manos (Ronchetti et al., 2016), una situación que no se "
           "repite cuando alguien usa la aplicación con las manos descubiertas.")
    cuerpo(doc, "Por ese motivo una parte del trabajo la hizo el equipo de estudiantes para pulir el sistema: "
           "probar cada seña frente a la cámara, registrar cuáles fallaban y, con esa información, ajustar el "
           "reconocimiento. Este proceso de prueba y ajuste permitió pasar de los datos públicos a un sistema que "
           "funcione en una situación real, con cámaras, iluminación y personas distintas de las de los videos de "
           "origen. Por ejemplo, el equipo registró que muchas palabras costaban detectarse. Al analizarlo se "
           "encontró la causa: en los videos de LSA64, en las señas que se hacen con una sola mano, la otra mano "
           "casi nunca aparece, mientras que frente a una computadora suele quedar a la vista, apoyada o a la "
           "altura de la cintura. En una prueba con personas que el modelo no vio al entrenar, con la otra mano "
           "en reposo a la vista, el modelo de palabras reconocía el 35 % de las señas y el de colores, el 16 %.")
    cuerpo(doc, "Para corregirlo, durante el entrenamiento se agregaron manos en reposo reales, tomadas de los "
           "videos del abecedario, y el reconocimiento en esa misma prueba subió al 92 % en las dos categorías. "
           "En el mismo proceso se ajustaron otros dos puntos. La seña «aceptar» se hace con una mano abierta "
           "que no se detectaba en los videos de LSA64; se completó con la forma de una palma tomada de una "
           "grabación propia. Además, los modelos se entrenan con la misma seña hecha a distintas velocidades y "
           "con pausas antes y después, como ocurre con alguien que recién aprende. La aplicación también deja "
           "que cada persona regule la confianza mínima para aceptar una seña, entre 30 % y 95 %, de modo de "
           "adaptarla a sus condiciones de uso.")

    doc.add_heading("Publicación y control de versiones", 2)
    cuerpo(doc, "El código se guarda en un repositorio de GitHub (MateoRo2003/LDS) y la aplicación se publica como "
           "sitio estático en un entorno «HTML5 · Vanilla Web» del campus del instituto. Como el reconocimiento "
           "ocurre en el navegador, no hace falta un servidor propio: alcanza con entregar los archivos del "
           "proyecto. Los navegadores solo permiten el acceso a la cámara en sitios con conexión segura (Mozilla, "
           "s.f.), y el campus la provee. La rama principal del repositorio contiene exactamente los 27 archivos "
           "que se suben al campus, junto con un archivo README.md que explica los pasos para subirlos; el código "
           "de entrenamiento, en Python, se guarda en otra rama del mismo repositorio. Como el campus conserva "
           "una copia de los archivos durante algunas horas, después de subir una versión nueva conviene "
           "recargar la página con Ctrl + Shift + R.")

    doc.add_heading("Alternativas y limitaciones", 2)
    cuerpo(doc, "Frente a otras opciones posibles, como una aplicación de escritorio o un servidor que reciba el "
           "video y devuelva el resultado, la solución elegida tiene tres ventajas: no exige instalar programas, "
           "no envía imágenes de las personas por internet y no requiere mantener un servidor. Sus limitaciones "
           "son cuatro. La primera es que el rendimiento depende de la computadora y del navegador de quien la "
           "usa, por lo que la guía recomienda Chrome o Edge, buena iluminación y mostrar la cara, los hombros y "
           "las manos. La segunda es que los modelos aprendieron de pocas personas y de videos con condiciones "
           "propias, por lo que pueden fallar con manos, ropa o iluminación muy distintas. La tercera es que los "
           "saludos y las preguntas se aprendieron de una sola persona cada grupo, por lo que su reconocimiento "
           "con otras personas todavía debe medirse. La cuarta es que el vocabulario se limita a 88 señas "
           "aisladas: la aplicación no interpreta frases.")


REQ_FUNCIONALES = [
    ("Acceso, cámara y detección", [
        ("RF1", "Abrir la aplicación desde el navegador",
         "Se accede por una dirección web, sin instalar programas ni crear cuentas."),
        ("RF2", "Solicitar permiso para usar la cámara",
         "Pide autorización al navegador y ofrece un botón para volver a intentar si no fue concedida."),
        ("RF3", "Seleccionar la cámara",
         "Si hay más de una cámara conectada, permite elegir cuál usar y recuerda la elección."),
        ("RF4", "Mostrar la imagen en vivo", "Muestra el video de la cámara en tiempo real, con la indicación «En vivo»."),
        ("RF5", "Detectar manos y cuerpo",
         "Ubica hasta dos manos (21 puntos cada una) y la nariz y los hombros como referencia."),
    ]),
    ("Reconocimiento de señas", [
        ("RF6", "Seleccionar una categoría",
         "Permite elegir abecedario, números, colores o palabras desde la barra lateral o desde «Explorar "
         "categorías»; en palabras, además, el grupo: básicas, saludos o preguntas."),
        ("RF7", "Reconocer posturas",
         "Identifica las letras y los números cuando la mano permanece quieta alrededor de medio segundo."),
        ("RF8", "Reconocer señas con movimiento",
         "Identifica los colores y las palabras de los tres grupos a partir de la seña completa, analizada como "
         "una secuencia de imágenes."),
        ("RF9", "Mostrar la seña detectada",
         "Muestra el nombre de la seña y su tipo (letra, número, color, palabra, saludo o pregunta)."),
        ("RF10", "Mostrar el nivel de confianza", "Indica con un porcentaje y con una barra qué tan segura está la detección."),
        ("RF11", "Descartar detecciones dudosas",
         "No acepta una seña si su confianza es menor al mínimo configurado o si no se sostiene en el tiempo."),
    ]),
    ("Práctica e historial", [
        ("RF12", "Seleccionar una seña para practicar",
         "Permite tocar una seña de la lista de la categoría o del grupo y activa el modo práctica."),
        ("RF13", "Indicar si la seña coincide",
         "Informa si lo realizado corresponde a la seña elegida o con cuál otra se la confundió."),
        ("RF14", "Registrar el historial de la sesión",
         "Guarda las últimas señas con su confianza y su hora; permite ver el historial completo (hasta 30 "
         "señas) y limpiarlo."),
    ]),
    ("Ajustes y estado del sistema", [
        ("RF15", "Ajustar la confianza mínima", "Deslizador entre 30 % y 95 %, con 60 % como valor inicial."),
        ("RF16", "Configurar la visualización",
         "Permite mostrar u ocultar los puntos de la mano sobre la imagen y ver la cámara como un espejo."),
        ("RF17", "Cambiar entre modo claro y oscuro", "Alterna el tema visual de la pantalla y recuerda la elección."),
        ("RF18", "Informar el estado del sistema",
         "Indica si está iniciando, activo o sin cámara, y da indicaciones como «Ponete de frente a la cámara»."),
        ("RF19", "Informar errores",
         "Explica cuándo falta el permiso, no hay cámara, la cámara está ocupada o no se pudo cargar el "
         "reconocimiento."),
        ("RF20", "Cargar los modelos por categoría",
         "Descarga el modelo de cada categoría y de cada grupo de palabras; los que no estén disponibles quedan "
         "fuera de servicio."),
    ]),
]

REQ_NO_FUNCIONALES = [
    ("RNF1", "Privacidad", "La imagen de la cámara debe procesarse en la computadora de la persona y no enviarse a "
     "ningún servidor."),
    ("RNF2", "Privacidad", "El navegador solo debe guardar preferencias (la cámara elegida y el tema visual), nunca "
     "datos personales ni imágenes."),
    ("RNF3", "Compatibilidad", "La aplicación debe funcionar en las versiones actuales de Chrome y de Edge."),
    ("RNF4", "Acceso", "No debe requerir instalar programas ni crear cuentas; se abre desde una dirección segura "
     "(HTTPS), necesaria para usar la cámara."),
    ("RNF5", "Rendimiento", "El reconocimiento debe responder en tiempo real, procesando cada imagen nueva de la "
     "cámara y usando la placa de video cuando esté disponible."),
    ("RNF6", "Rendimiento", "La primera carga puede demorar, pero los ingresos posteriores deben abrir rápido."),
    ("RNF7", "Usabilidad", "La interfaz debe ser clara y estar en español, con acceso directo a las categorías, a la "
     "seña detectada y al estado del sistema, y con modos claro y oscuro."),
    ("RNF8", "Robustez", "Los errores de cámara o de carga deben informarse en pantalla sin cerrar la aplicación; si "
     "falta un modelo, solo esa categoría o ese grupo queda fuera de servicio."),
    ("RNF9", "Fiabilidad", "El reconocimiento debe ser independiente de la distancia y la posición de la persona "
     "frente a la cámara, valer con cualquiera de las dos manos y no verse afectado si la mano que no participa "
     "de la seña queda quieta a la vista."),
    ("RNF10", "Mantenibilidad", "La estructura debe ser modular, con la pantalla, el control de la cámara, el "
     "reconocimiento y los modelos separados; una categoría o un grupo nuevo se agrega sumando un modelo."),
    ("RNF11", "Portabilidad", "Debe poder publicarse como sitio web estático, sin servidor propio: 27 archivos que "
     "ocupan alrededor de 3,6 MB."),
    ("RNF12", "Almacenamiento", "Los modelos deben estar separados por categoría y por grupo, y los de mayor tamaño "
     "divididos en partes de unos 250 KB."),
    ("RNF13", "Legalidad", "El uso de datos públicos debe respetar sus licencias (CC BY-NC-SA en el caso de LSA64) y "
     "citar sus fuentes, por lo que el proyecto es educativo y no comercial."),
]


def requerimientos(doc):
    doc.add_heading("Requerimientos", 1)
    cuerpo(doc, "Los requerimientos se relevaron a partir del funcionamiento de la aplicación y de los problemas "
           "identificados en la investigación. Se dividen en funcionales, que describen lo que el sistema permite "
           "hacer, y no funcionales, que describen condiciones de calidad, seguridad y desempeño.")

    doc.add_heading("Requerimientos funcionales (RF)", 2)
    cuerpo(doc, "Detalles de las funciones que el sistema permite realizar a la persona que lo usa, agrupadas según "
           "la parte de la aplicación a la que corresponden.")
    tabla = doc.add_table(rows=0, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(tabla)
    for grupo, filas in REQ_FUNCIONALES:
        celdas = tabla.add_row().cells
        unida = celdas[0].merge(celdas[2])
        texto_celda(unida, grupo, 10, negrita=True, color=BORDO)
        sombrear(unida, ROSA_GRUPO)
        for i, (rid, nombre, detalle) in enumerate(filas):
            celdas = tabla.add_row().cells
            texto_celda(celdas[0], rid, 10, negrita=True, color=BORDO)
            texto_celda(celdas[1], nombre, 10)
            texto_celda(celdas[2], detalle, 10)
            if i % 2:
                for c in celdas:
                    sombrear(c, ROSA_FILA)
    anchos(tabla, [1.5, 4.6, 9.9])

    doc.add_heading("Requerimientos no funcionales (RNF)", 2)
    cuerpo(doc, "En esta parte se detallan las condiciones de calidad, seguridad y desempeño que el sistema debe "
           "cumplir, identificadas según el atributo de calidad al que responden.")
    tabla = doc.add_table(rows=1, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(tabla)
    for celda, texto in zip(tabla.rows[0].cells, ("ID", "Atributo", "Requerimiento")):
        texto_celda(celda, texto, 10, negrita=True, color=BORDO)
        sombrear(celda, ROSA_GRUPO)
    encabezado = tabla.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    encabezado.append(repetir)
    for i, (rid, atributo, detalle) in enumerate(REQ_NO_FUNCIONALES):
        celdas = tabla.add_row().cells
        texto_celda(celdas[0], rid, 10, negrita=True, color=BORDO)
        texto_celda(celdas[1], atributo, 10)
        texto_celda(celdas[2], detalle, 10)
        if i % 2:
            for c in celdas:
                sombrear(c, ROSA_FILA)
    anchos(tabla, [1.6, 3.0, 11.4])


def diseno(doc):
    doc.add_heading("Diseño de la solución", 1)
    cuerpo(doc, "Como SIGNALIS está desarrollada en JavaScript sin frameworks, no se modeló con un diagrama de "
           "clases. En su lugar, esta sección presenta la arquitectura del sistema y el flujo de uso de la "
           "aplicación.")
    doc.add_heading("Arquitectura del sistema", 2)
    cuerpo(doc, "La Figura muestra los componentes de SIGNALIS. Todo el procesamiento ocurre dentro del navegador: "
           "la cámara entrega la imagen a la aplicación, MediaPipe ubica las manos, el motor de reconocimiento "
           "identifica la seña con ayuda de los modelos de cada categoría y la interfaz muestra el resultado. El "
           "sitio alojado en el entorno del campus solo entrega los archivos del proyecto.")
    figura(doc, "arquitectura.png", 15.5)
    cuerpo(doc, "El sitio consta de 27 archivos. En la raíz están index.html y style.css, que forman la pantalla; "
           "app.js, que maneja la cámara y los botones; y motor.js, que reconoce las señas. La carpeta modelos "
           "contiene 22 archivos JSON: un modelo por categoría y por grupo de palabras (abecedario, números, "
           "colores, palabras básicas, saludos y preguntas), más las cuatro partes en que se dividen los de "
           "colores, palabras, saludos y preguntas. La carpeta vendor/tasks-vision contiene vision_bundle.js, la "
           "librería MediaPipe Tasks Vision. A estos se suma el archivo README.md, con las instrucciones para "
           "subir y usar el proyecto, que no se sube al entorno del campus. Los nombres de archivos y carpetas "
           "deben respetarse tal como figuran, en minúsculas, porque la aplicación los busca por esa ruta.")
    doc.add_heading("Flujo de uso", 2)
    cuerpo(doc, "Resumen de cómo se usa la aplicación. Al abrirla, carga los modelos y los detectores de MediaPipe "
           "y pide permiso para usar la cámara. La persona elige una categoría (en palabras, también el grupo) y, "
           "si quiere practicar, una seña de la lista; se ubica de frente a la cámara y hace la seña. Por cada "
           "imagen nueva, el sistema ubica las manos y el cuerpo y el motor calcula qué seña es la más probable. "
           "Si la confianza supera el mínimo configurado (60 % por defecto), la aplicación muestra la seña con su "
           "nivel de confianza y la agrega al historial; en el modo práctica, además, indica si coincide con la "
           "seña elegida.")
    figura(doc, "flujo_de_uso.png", 14.5)
    cuerpo(doc, "Es importante distinguir dos tipos de señas. Las letras y los números son posturas, por lo que se "
           "debe mantener la mano quieta medio segundo. Los colores y las palabras son movimientos, por lo que se "
           "hace la seña completa una vez y se baja la mano; al bajarla, el sistema da por terminada la seña y "
           "empieza de cero con la siguiente.")


REFERENCIAS = [
    "Agencia Nacional de Discapacidad. (s.f.). *Interpretación en lengua de señas argentina (LSA)*. Educ.ar. "
    "https://www.educ.ar/recursos/158988/interpretacion-en-lengua-de-senas-argentina-lsa",
    "chany3h. (s.f.). *Lengua de señas Argentina LSA Clase 1 ABCedario* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=YKHsz05061I",
    "Chollet, F. (2015). *Keras*. https://keras.io",
    "Dirección de Documentación. (s.f.-a). *Alfabeto dactilológico LSA* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=bhzaMJw_NbE",
    "Dirección de Documentación. (s.f.-b). *Números en LSA* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=gLSSENEOGD4",
    "Dirección de Documentación. (s.f.-c). *Organizadores de preguntas en LSA* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=x57rpA804yA",
    "Dirección de Documentación. (s.f.-d). *Saludos y fórmulas de cortesía en LSA* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=g18JlOMg5Lk",
    "Google. (s.f.). *MediaPipe: On-device machine learning for everyone*. GitHub. "
    "https://github.com/google-ai-edge/mediapipe",
    "Hochreiter, S. y Schmidhuber, J. (1997). Long short-term memory. *Neural Computation, 9*(8), 1735-1780. "
    "https://doi.org/10.1162/neco.1997.9.8.1735",
    "Ley 27.710 de 2023. Lengua de Señas Argentina. 13 de abril de 2023. Honorable Congreso de la Nación "
    "Argentina. https://www.argentina.gob.ar/normativa/nacional/ley-27710-383041/normas-modifican",
    "LSAorg. (s.f.). *Impulsar*. GitHub. https://github.com/LSAorg/impulsar",
    "Lucho-code. (s.f.). *Entreclases: App para aprender y entender lenguaje de señas*. GitHub. "
    "https://github.com/Lucho-code/entreclases",
    "Mohammedalmohawis. (s.f.). *lsa-sign-language-to-speech: End-to-end LSA sign language to speech AI*. GitHub. "
    "https://github.com/mohammedalmohawis/lsa-sign-language-to-speech",
    "Mozilla. (s.f.). *MediaDevices: getUserMedia() method*. MDN Web Docs. "
    "https://developer.mozilla.org/en-US/docs/Web/API/MediaDevices/getUserMedia",
    "Ronchetti, F., Quiroga, F. M., Estrebou, C., Lanzarini, L. y Rosete, A. (2016). LSA64: An Argentinian sign "
    "language dataset. En *XXII Congreso Argentino de Ciencias de la Computación (CACIC 2016)*. Universidad "
    "Nacional de La Plata. https://sedici.unlp.edu.ar/handle/10915/56764",
    "Sarria, C. (s.f.). *Números en LSA - Lengua de Señas Argentina* [Video]. YouTube. "
    "https://www.youtube.com/watch?v=9ghx0rmFKKg",
    "Schuster, M. y Paliwal, K. K. (1997). Bidirectional recurrent neural networks. *IEEE Transactions on Signal "
    "Processing, 45*(11), 2673-2681. https://doi.org/10.1109/78.650093",
    "Servicio de Información sobre Discapacidad. (s.f.). *¿Cuántos tipos de lengua de signos existen en el "
    "mundo?* Universidad de Salamanca. https://sid-inico.usal.es/noticias/260869/",
    "Universidad de Buenos Aires, Facultad de Filosofía y Letras. (s.f.). *Lengua de señas argentina*. "
    "https://www.idiomas.filo.uba.ar/es/lengua-de-senas-argentina",
    "Universidad Popular, Dirección de Inclusión Educativa. (s.f.). *Abecedario LSA* [Video]. Facebook. "
    "https://www.facebook.com/watch/?v=265056705412909",
]


def referencias(doc):
    salto_de_pagina(doc)
    doc.add_heading("Referencias bibliográficas", 1)
    for ref in REFERENCIAS:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        agregar_texto(p, ref)


def titulo_anexo(doc, texto, nivel):
    """Titulos de la guia dentro del anexo: con el color de la guia, sin entrar al indice."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if nivel == 0 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto)
    fuente(r, 16 if nivel == 0 else 12, negrita=True, color=BORDO)
    r.italic = nivel == 2
    return p


def anexo(doc):
    salto_de_pagina(doc)
    doc.add_heading("Anexo", 1)
    guia.escribir_guia(doc, titulo=titulo_anexo, ancho_util=ANCHO_UTIL, compacta=True)


def exportar_pdf(ruta_docx):
    """Abre el .docx en Word (por PowerShell), completa el indice y lo guarda tambien como PDF."""
    ruta_docx = os.path.abspath(ruta_docx)
    ruta_pdf = os.path.splitext(ruta_docx)[0] + ".pdf"
    script = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open('{ruta_docx}')
  # estilos del indice (toc 1 = -20, toc 2 = -21): interlineado 1,5 para que entre en una hoja
  foreach ($n in -20, -21) {{ $doc.Styles.Item($n).ParagraphFormat.LineSpacingRule = 1; $doc.Styles.Item($n).ParagraphFormat.SpaceAfter = 0 }}
  foreach ($toc in $doc.TablesOfContents) {{ $toc.Update() }}
  $doc.Fields.Update() | Out-Null
  foreach ($toc in $doc.TablesOfContents) {{ $toc.Update() }}
  $doc.Save()
  $doc.SaveAs2('{ruta_pdf}', 17)
  $doc.Close($false)
}} finally {{ $word.Quit() }}
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return ruta_pdf


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf", action="store_true", help="completar el indice con Word y exportar a PDF")
    args = parser.parse_args()

    doc = Document()
    preparar_estilos(doc)
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = Cm(21.0), Cm(29.7)
    seccion.left_margin = seccion.right_margin = Cm(2.5)
    seccion.top_margin = seccion.bottom_margin = Cm(2.5)

    portada(doc)
    indice(doc)
    seccion_con_numeros(doc)
    descripcion(doc)
    objetivos(doc)
    investigacion(doc)
    tecnologias(doc)
    requerimientos(doc)
    diseno(doc)
    referencias(doc)
    anexo(doc)
    compactar_tablas(doc)

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    doc.save(SALIDA)
    print(f"{SALIDA}: {os.path.getsize(SALIDA) / 1e6:.1f} MB")
    if args.pdf:
        print(exportar_pdf(SALIDA))


if __name__ == "__main__":
    main()
